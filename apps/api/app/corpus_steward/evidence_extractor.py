"""Deterministic, source-anchored extraction for DAK PDF and XLSX assets."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal

import pymupdf
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from app.schemas.corpus import LocatorKind, SourceAnchor

PDF_MEDIA_TYPE = "application/pdf"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class EvidenceExtractionError(ValueError):
    def __init__(self, reason_code: str, details: str) -> None:
        self.reason_code = reason_code
        super().__init__(details)


@dataclass(frozen=True)
class ExtractedEvidenceUnit:
    source_unit_id: str
    content_exact: str
    content_search: str
    anchors: tuple[SourceAnchor, ...]


@dataclass(frozen=True)
class ExtractedAsset:
    expected_source_units: int
    empty_source_units: int
    units: tuple[ExtractedEvidenceUnit, ...]
    # Narrative extraction removes running headers and footers. Counted here so the drop
    # is visible to coverage accounting rather than silent. Always 0 on the DAK path,
    # which removes nothing.
    dropped_boilerplate_blocks: int = 0


class DAKSourceExtractor:
    def extract(self, content: bytes, *, media_type: str, source_uri: str) -> ExtractedAsset:
        if media_type == PDF_MEDIA_TYPE:
            return self._pdf(content, source_uri=source_uri)
        if media_type == XLSX_MEDIA_TYPE:
            return self._xlsx(content, source_uri=source_uri)
        raise EvidenceExtractionError(
            "UNSUPPORTED_NARRATIVE_MEDIA_TYPE",
            f"materialization does not support {media_type}",
        )

    @staticmethod
    def _pdf(content: bytes, *, source_uri: str) -> ExtractedAsset:
        try:
            document = pymupdf.open(stream=content, filetype="pdf")
        except Exception as error:  # PyMuPDF exposes several parser exception classes.
            raise EvidenceExtractionError("PDF_OPEN_FAILED", str(error)) from error
        try:
            if document.page_count < 1:
                raise EvidenceExtractionError("PDF_HAS_NO_PAGES", "PDF contains no pages")
            units: list[ExtractedEvidenceUnit] = []
            empty = 0
            for page_index in range(document.page_count):
                page = document.load_page(page_index)
                blocks = []
                for block in page.get_text("blocks", sort=True):
                    text = str(block[4]).strip()
                    if not text or int(block[6]) != 0:
                        continue
                    blocks.append((block[:4], text))
                if not blocks:
                    empty += 1
                    continue
                exact = "\n\n".join(text for _, text in blocks)
                units.append(
                    ExtractedEvidenceUnit(
                        source_unit_id=f"pdf:page:{page_index + 1}",
                        content_exact=exact,
                        content_search=_search_view(exact),
                        anchors=tuple(
                            SourceAnchor(
                                kind=LocatorKind.PDF,
                                source_uri=source_uri,
                                pdf_page=page_index + 1,
                                bbox=tuple(float(value) for value in bbox),
                            )
                            for bbox, _ in blocks
                        ),
                    )
                )
            return ExtractedAsset(
                expected_source_units=document.page_count,
                empty_source_units=empty,
                units=tuple(units),
            )
        finally:
            document.close()

    @staticmethod
    def _xlsx(content: bytes, *, source_uri: str) -> ExtractedAsset:
        try:
            workbook = load_workbook(
                io.BytesIO(content),
                read_only=True,
                data_only=False,
                keep_links=False,
            )
        except Exception as error:
            raise EvidenceExtractionError("XLSX_OPEN_FAILED", str(error)) from error
        try:
            if not workbook.sheetnames:
                raise EvidenceExtractionError("XLSX_HAS_NO_WORKSHEETS", "workbook is empty")
            units: list[ExtractedEvidenceUnit] = []
            empty = 0
            expected = 0
            for worksheet in workbook.worksheets:
                max_row = max(1, worksheet.max_row or 1)
                expected += max_row
                for row_index, row in enumerate(
                    worksheet.iter_rows(min_row=1, max_row=max_row), start=1
                ):
                    cells = [
                        (column_index, _cell_text(cell.value))
                        for column_index, cell in enumerate(row, start=1)
                        if cell.value is not None and _cell_text(cell.value) != ""
                    ]
                    if not cells:
                        empty += 1
                        continue
                    exact = "\n".join(
                        f"{get_column_letter(column_index)}{row_index}={value}"
                        for column_index, value in cells
                    )
                    units.append(
                        ExtractedEvidenceUnit(
                            source_unit_id=f"xlsx:{worksheet.title}:row:{row_index}",
                            content_exact=exact,
                            content_search=_search_view(" ".join(value for _, value in cells)),
                            anchors=tuple(
                                SourceAnchor(
                                    kind=LocatorKind.TABLE_CELL,
                                    source_uri=source_uri,
                                    table_id=worksheet.title,
                                    row_index=row_index - 1,
                                    column_index=column_index - 1,
                                )
                                for column_index, _ in cells
                            ),
                        )
                    )
            return ExtractedAsset(
                expected_source_units=expected,
                empty_source_units=empty,
                units=tuple(units),
            )
        finally:
            workbook.close()


def _cell_text(value: object) -> str:
    if isinstance(value, ArrayFormula):
        return value.text or f"ARRAY_FORMULA:{value.ref}"
    if isinstance(value, DataTableFormula):
        return "DATA_TABLE_FORMULA:" + ";".join(f"{key}={item}" for key, item in value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value).strip()


def _search_view(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()
